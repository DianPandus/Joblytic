"""Ekstraksi CV ke profil terstruktur.

Alur: baca teks (PDF/DOCX) -> hapus identitas -> LLM mengisi skema Pydantic ->
buang entri yang tidak bisa ditemukan di teks CV (LLM tidak boleh mengarang).
"""

import io
import re
import zipfile
from dataclasses import dataclass
from typing import Literal

from pydantic import BaseModel, Field

from app.grounding import clean, supported, words
from app.llm import LLMClient, LLMError

PROMPT_VERSION = "cv-extract-v1"
MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_TEXT_CHARS = 20_000

PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

Section = Literal[
    "education", "experience", "organization", "project", "skill", "certification", "language"
]


class CVError(Exception):
    """Masalah pada berkas CV yang perlu diperbaiki pengguna (pesan aman ditampilkan)."""


# --- Skema keluaran LLM -------------------------------------------------------


class ExtractedEntry(BaseModel):
    section: Section
    title: str = Field(
        description="Nama utama entri, disalin dari CV: gelar/jurusan, jabatan, nama proyek, "
        "nama skill, nama sertifikat, atau nama bahasa."
    )
    organization: str | None = Field(
        description="Institusi, perusahaan, organisasi, atau penerbit. null jika tidak ada."
    )
    location: str | None = Field(description="Kota/negara jika tertulis. null jika tidak ada.")
    start_month: str | None = Field(
        description="Format YYYY-MM. Jika hanya tahun, pakai YYYY-01. null jika tidak ada."
    )
    end_month: str | None = Field(
        description="Format YYYY-MM. null jika tidak ada atau masih berlangsung."
    )
    is_current: bool = Field(description="true jika tertulis masih berlangsung/sekarang.")
    description: str | None = Field(
        description="Keterangan singkat yang tertulis: IPK, tingkat kemahiran, kategori skill. "
        "null jika tidak ada."
    )
    bullets: list[str] = Field(
        description="Butir pencapaian/tanggung jawab, disalin apa adanya dari CV. "
        "Kosongkan jika tidak ada."
    )


class CVExtraction(BaseModel):
    entries: list[ExtractedEntry]


SYSTEM_PROMPT = """Kamu mengekstrak isi CV ke data terstruktur.
Aturan:
- Hanya salin informasi yang benar-benar tertulis di CV. Jangan menebak, menyimpulkan,
  menerjemahkan, atau menambah apa pun.
- Field yang tidak tertulis diisi null (atau list kosong untuk bullets).
- Salin teks apa adanya, dalam bahasa aslinya.
- Satu entri per pendidikan, pengalaman kerja/magang, organisasi, proyek, sertifikat,
  dan bahasa. Setiap skill menjadi satu entri section "skill" dengan title nama skill.
- Abaikan identitas pribadi (nama, kontak, alamat); teks seperti [NAMA] atau [EMAIL]
  adalah hasil penyensoran dan jangan dimasukkan."""

STRICT_SUFFIX = """
PENTING: percobaan sebelumnya tidak valid. Keluarkan JSON yang persis mengikuti skema,
dengan section hanya salah satu nilai yang diizinkan dan tanggal berformat YYYY-MM."""


# --- Membaca berkas -----------------------------------------------------------


def detect_type(data: bytes) -> str:
    """Tentukan tipe dari isi berkas, bukan dari nama atau header yang dikirim klien."""
    if data.startswith(b"%PDF-"):
        return PDF
    if data.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                if "word/document.xml" in z.namelist():
                    return DOCX
        except zipfile.BadZipFile:
            pass
    raise CVError("Format berkas tidak didukung. Unggah CV dalam bentuk PDF atau DOCX.")


def extract_text(data: bytes, mime: str) -> str:
    try:
        if mime == PDF:
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                raise CVError("PDF terkunci password. Unggah versi tanpa password.")
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        else:
            from docx import Document

            doc = Document(io.BytesIO(data))
            lines = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    lines.append(" | ".join(cell.text for cell in row.cells))
            text = "\n".join(lines)
    except CVError:
        raise
    except Exception as exc:  # parser pihak ketiga bisa melempar apa saja untuk berkas rusak
        raise CVError("Berkas CV tidak bisa dibaca. Pastikan berkasnya tidak rusak.") from exc

    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n", text).strip()
    if len(text) < 50:
        raise CVError(
            "Teks CV tidak terbaca. CV hasil scan atau berupa gambar belum didukung; "
            "unggah CV yang teksnya bisa diseleksi."
        )
    return text[:MAX_TEXT_CHARS]


# --- Penghapusan identitas sebelum dikirim ke LLM (PRD: privasi) ----------------

EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+(\.[\w-]+)+")
URL_RE = re.compile(r"(https?://|www\.)\S+|\b(linkedin|github)\.com/\S+", re.IGNORECASE)
PHONE_RE = re.compile(r"(?<![\w])(\+?62|0)\s?8[\d\s().-]{7,14}\d")
ADDRESS_RE = re.compile(
    r"^.*(\balamat\b|\baddress\b|\bjl\.|\bjln\b|\bjalan\b|\brt\s*\d|\brw\s*\d|\bkec\.|"
    r"\bkel\.|\bkode pos\b).*$",
    re.IGNORECASE | re.MULTILINE,
)
NAME_LINE_RE = re.compile(r"^[A-Za-z][A-Za-z.' ]{2,60}$")


def redact(text: str, full_name: str | None) -> str:
    text = EMAIL_RE.sub("[EMAIL]", text)
    text = URL_RE.sub("[URL]", text)
    text = PHONE_RE.sub("[TELEPON]", text)
    text = ADDRESS_RE.sub("[ALAMAT]", text)

    # Nama: dari profil akun, ditambah baris pertama CV bila bentuknya seperti nama.
    names: set[str] = set()
    if full_name:
        names.add(full_name)
    first_line = next((line.strip() for line in text.splitlines() if line.strip()), "")
    if NAME_LINE_RE.match(first_line) and 2 <= len(first_line.split()) <= 5:
        names.add(first_line)
    tokens = {t for name in names for t in name.split() if len(t) >= 3}
    for token in sorted(tokens, key=len, reverse=True):
        text = re.sub(rf"\b{re.escape(token)}\b", "[NAMA]", text, flags=re.IGNORECASE)
    return text


# --- Grounding: entri harus bisa ditemukan di teks CV ---------------------------


def _month(value: str | None) -> str | None:
    if not value:
        return None
    m = re.fullmatch(r"(\d{4})-(\d{2})", value.strip())
    if not m or not (1 <= int(m.group(2)) <= 12) or not (1950 <= int(m.group(1)) <= 2100):
        return None
    return f"{m.group(1)}-{m.group(2)}-01"


@dataclass
class ExtractionResult:
    entries: list[dict]
    dropped_entries: int
    dropped_bullets: int


def ground(extraction: CVExtraction, source_text: str) -> ExtractionResult:
    source_words = set(words(source_text))
    entries: list[dict] = []
    dropped_entries = dropped_bullets = 0

    for e in extraction.entries:
        title = clean(e.title, 300)
        if not title or not supported(title, source_words):
            dropped_entries += 1
            continue
        bullets = []
        for b in e.bullets:
            b = clean(b, 1000)
            if b and supported(b, source_words):
                bullets.append(b)
            elif b:
                dropped_bullets += 1

        def keep(value: str | None, limit: int) -> str | None:
            value = clean(value, limit)
            return value if value and supported(value, source_words) else None

        start, end = _month(e.start_month), _month(e.end_month)
        if start and end and end < start:
            end = None
        entries.append(
            {
                "section": e.section,
                "title": title,
                "organization": keep(e.organization, 300),
                "location": keep(e.location, 200),
                "start_month": start,
                "end_month": None if e.is_current else end,
                "is_current": e.is_current,
                "description": keep(e.description, 2000),
                "bullets": bullets,
            }
        )
    return ExtractionResult(entries, dropped_entries, dropped_bullets)


async def extract_profile(llm: LLMClient, redacted_text: str) -> ExtractionResult:
    """Panggil LLM (maksimal dua kali, kedua dengan instruksi lebih ketat), lalu grounding."""
    prompt = f"Teks CV:\n<<<\n{redacted_text}\n>>>"
    last_error: LLMError | None = None
    for system in (SYSTEM_PROMPT, SYSTEM_PROMPT + STRICT_SUFFIX):
        try:
            extraction = await llm.generate_json(system, prompt, CVExtraction)
            return ground(extraction, redacted_text)
        except LLMError as exc:
            if exc.__class__ is not LLMError:  # kuota/konfigurasi: jangan diulang
                raise
            last_error = exc
    assert last_error is not None
    raise last_error
